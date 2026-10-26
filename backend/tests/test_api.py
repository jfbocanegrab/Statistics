import base64
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import create_app

SAMPLES = Path(__file__).resolve().parents[2] / "samples"

SIMPLE = b"""timestamp,cpu,trafico
2026-10-14 00:00,10,100
2026-10-14 12:00,20,200
2026-10-22 00:00,30,300
2026-10-22 12:00,50,500
"""

EURO = "Fecha;Hora;Latencia (ms)\n14/10/2026;00:00;10,5\n22/10/2026;00:00;21,0\n22/10/2026;23:45;1.000,0\n".encode("cp1252")

# Reportes del gestor: mismo formato, distinto grupo de KPI, mismas fechas y un KPI repetido
TRAFFIC = b"""

5G Performance Report
Save Time :2026-10-26 10:53:50

Date,Whole Network,Integrity,NetSF Average User Number,NetSF DL Traffic Volume GB(GB)
2026-10-01,Whole Network,98%,100,1000
2026-10-02,Whole Network,97%,110,1100
2026-10-03,Whole Network,98%,120,1200
Total 3 Records
"""
THROUGHPUT = b"""5G Performance Report
Save Time :2026-10-26 10:54:10

Date,Whole Network,Integrity,NetSF Average User Number,NetSF Average Cell DL throughput Mbps(Mbit/s)
2026-10-01,Whole Network,98%,100,55.5
2026-10-02,Whole Network,97%,110,56.5
2026-10-03,Whole Network,98%,120,57.5
Total 3 Records
"""


@pytest.fixture()
def client():
    return create_app().test_client()


def post(client, url, data, cfg):
    files = data if isinstance(data, list) else [("x.csv", data)]
    return client.post(url, data={"file": [(io.BytesIO(d), n) for n, d in files], "config": json.dumps(cfg)},
                       content_type="multipart/form-data")


def test_analyze_and_compare_days(client):
    cfg = {"executions": [{"id": "e1", "label": "Change", "date": "2026-10-20", "mode": "day"}],
           "compare_days": ["2026-10-14", "2026-10-22"]}
    r = post(client, "/api/analyze", SIMPLE, cfg)
    assert r.status_code == 200, r.json
    j = r.json
    assert j["metrics"] == ["NetSF cpu", "NetSF trafico"]                 # prefijo NetSF
    assert "group_column" not in j                                        # sin grupos: un solo reporte
    assert len(j["charts"][0]["series"]) == 1
    c = j["comparison"]
    assert (c["day1"], c["day2"]) == ("2026-10-14", "2026-10-22")        # independientes de la ejecución
    cpu = next(x for x in c["rows"] if x["metric"] == "NetSF cpu")
    assert "group" not in cpu
    assert (cpu["v1"], cpu["v2"], cpu["delta"], cpu["delta_pct"]) == (15.0, 40.0, 25.0, pytest.approx(166.666667))


def test_no_compare_days_no_table(client):
    j = post(client, "/api/analyze", SIMPLE, {"executions": [{"date": "2026-10-22", "mode": "day"}]}).json
    assert j["comparison"] is None


def test_compare_days_need_two(client):
    r = post(client, "/api/analyze", SIMPLE, {"compare_days": ["2026-10-14"]})
    assert r.status_code == 400 and "both comparison days" in r.json["error"]


def test_euro_format_separate_date_time(client):
    cfg = {"compare_days": ["2026-10-14", "2026-10-22"], "agg": "max"}
    j = post(client, "/api/analyze", EURO, cfg).json
    assert j["metrics"] == ["NetSF Latencia (ms)"]
    assert "Fecha + Hora" == j["timestamp_column"]
    row = j["comparison"]["rows"][0]
    assert row["v2"] == 1000.0 and row["v1"] == 10.5


def test_metric_filter(client):
    j = post(client, "/api/analyze", SIMPLE, {"metrics": ["NetSF trafico"]}).json
    assert [c["metric"] for c in j["charts"]] == ["NetSF trafico"]


def test_repeated_dates_in_one_file_are_averaged(client):
    data = b"Date,KPI\n2026-10-01,10\n2026-10-01,20\n2026-10-02,5\n"
    j = post(client, "/api/analyze", data, {}).json
    assert j["summary"]["rows"] == 2
    assert j["charts"][0]["series"][0]["points"][0][1] == 15.0


def test_reports(client):
    cfg = {"executions": [{"date": "2026-10-22", "mode": "datetime", "time": "06:30", "label": "Window 1"},
                          {"date": "2026-10-22", "mode": "day", "label": "Window 2"}], "title": "Test",
           "compare_days": ["2026-10-14", "2026-10-22"]}
    pdf = post(client, "/api/report/pdf", SIMPLE, cfg)
    assert pdf.status_code == 200 and pdf.data[:4] == b"%PDF"
    x = post(client, "/api/report/xlsx", SIMPLE, cfg)
    assert x.status_code == 200 and x.data[:2] == b"PK"


@pytest.mark.parametrize("data,msg", [
    (b"", "empty"),
    (b"a,b\nx,y\nz,w\n", "date/time"),
    (b"fecha,nombre\n2026-01-01,x\n", "numeric"),
])
def test_errors(client, data, msg):
    r = post(client, "/api/analyze", data, {})
    assert r.status_code == 400 and msg in r.json["error"]


def test_bad_execution_date(client):
    r = post(client, "/api/analyze", SIMPLE, {"executions": [{"date": "22/10/2026"}]})
    assert r.status_code == 400


def test_iso_dates_not_dayfirst(client):
    data = b"timestamp,v\n2026-10-06 00:00:00,1\n2026-10-06 00:15:00,2\n"
    j = post(client, "/api/analyze", data, {}).json
    assert j["summary"]["start"].startswith("2026-10-06") and j["summary"]["step_minutes"] == 15.0


def test_dayfirst_slash_dates(client):
    data = b"fecha,v\n05/10/2026,1\n06/10/2026,2\n"
    j = post(client, "/api/analyze", data, {}).json
    assert j["summary"]["start"].startswith("2026-10-05")


def test_preview_pages(client):
    cfg = {"executions": [{"date": "2026-10-22", "mode": "day"}], "compare_days": ["2026-10-14", "2026-10-22"]}
    j = post(client, "/api/report/preview", SIMPLE, cfg).json
    assert base64.b64decode(j["pdf"])[:4] == b"%PDF"
    assert len(j["pages"]) >= 2 and base64.b64decode(j["pages"][0])[:4] == b"\x89PNG"
    assert j["filename"].endswith(".pdf")


def test_real_export_with_title_lines_footer_and_nil(client):
    """Formato real de exportación (5G Performance Report): título antes, pie 'Total N Records', NIL."""
    data = (SAMPLES / "ejemplo_5g_traffic.csv").read_bytes()
    cfg = {"executions": [{"date": "2026-10-22", "mode": "day"}], "compare_days": ["2026-10-17", "2026-10-25"]}
    r = post(client, "/api/analyze", data, cfg)
    assert r.status_code == 200, r.json
    j = r.json
    assert j["timestamp_column"] == "Date"
    assert len(j["metrics"]) == 9 and all(m.startswith("NetSF") for m in j["metrics"])
    assert "NetSF Average Uplink PDCP Packet Loss Rate(%)" in j["metrics"]
    assert "NetSF DL Traffic Volume GB(GB)" in j["metrics"]                 # no duplica el prefijo
    assert j["summary"]["rows"] == 30 and j["summary"]["step_minutes"] == 1440
    row = next(x for x in j["comparison"]["rows"] if x["metric"] == "NetSF Maximum User Number")
    assert (row["v1"], row["v2"]) == (996.0, 1158.0)
    assert post(client, "/api/report/pdf", data, cfg).status_code == 200
    assert post(client, "/api/report/xlsx", data, cfg).status_code == 200


# ---------------------------------------------------------------- varios CSV = un solo reporte
def test_multiple_files_same_dates_merge_by_date_and_dedupe_kpis(client):
    files = [("Report(Traffic).csv", TRAFFIC), ("Report(Throughput).csv", THROUGHPUT)]
    j = post(client, "/api/analyze", files, {"compare_days": ["2026-10-01", "2026-10-03"]}).json
    assert j["summary"]["rows"] == 3                                      # una fila por fecha, no 6
    # KPI repetidos una sola vez, en orden de prioridad: Traffic, Users, ..., Other KPIs
    assert j["metrics"] == ["NetSF DL Traffic Volume GB(GB)", "NetSF Average User Number",
                            "NetSF Average Cell DL throughput Mbps(Mbit/s)", "NetSF Integrity"]
    assert all(len(c["series"]) == 1 and len(c["series"][0]["points"]) == 3 for c in j["charts"])
    assert any("taken only once" in w for w in j["warnings"])
    assert not any("different columns" in w for w in j["warnings"])
    assert [f["new_kpis"] for f in j["files"]] == [3, 1]
    row = next(x for x in j["comparison"]["rows"] if x["metric"] == "NetSF Average Cell DL throughput Mbps(Mbit/s)")
    assert (row["v1"], row["v2"]) == (55.5, 57.5)


def test_multiple_files_consecutive_periods_are_joined(client):
    part1 = b"Date,KPI A\n2026-10-01,1\n2026-10-02,2\n"
    part2 = b"Date,KPI A\n2026-10-03,3\n2026-10-04,4\n"
    j = post(client, "/api/analyze", [("p1.csv", part1), ("p2.csv", part2)], {}).json
    assert j["summary"]["rows"] == 4
    assert [p[1] for p in j["charts"][0]["series"][0]["points"]] == [1, 2, 3, 4]


def test_duplicated_kpi_keeps_first_file_value(client):
    a = b"Date,KPI A\n2026-10-01,1\n"
    b = b"Date,KPI A\n2026-10-01,999\n2026-10-02,2\n"
    j = post(client, "/api/analyze", [("a.csv", a), ("b.csv", b)], {}).json
    assert [p[1] for p in j["charts"][0]["series"][0]["points"]] == [1, 2]   # la fecha nueva sí se agrega


def test_auto_status_and_colors_with_threshold_and_direction():
    from app.analysis import auto_status, higher_is_better
    assert higher_is_better("NetSF DL Traffic Volume GB(GB)")
    assert not higher_is_better("NetSF Average Uplink PDCP Packet Loss Rate(%)")
    assert not higher_is_better("NetSF SgNB Abnormal Release Rate (CU)(%)")
    assert not higher_is_better("NetSF SgNB-Trig SgNB Abnorm Rel Rate(%)")           # nombre abreviado
    assert not higher_is_better("NetSF UL Interference (dBm)(dBm)")
    assert auto_status("NetSF Traffic", 3.0) == "STABLE" and auto_status("NetSF Traffic", -3.0) == "STABLE"
    assert auto_status("NetSF Traffic", 3.01) == "IMPROVED" and auto_status("NetSF Traffic", -5) == "DEGRADED"
    assert auto_status("NetSF Packet Loss Rate", 10) == "DEGRADED" and auto_status("NetSF Packet Loss Rate", -10) == "IMPROVED"
    assert auto_status("NetSF Traffic", None) is None


def test_suggested_status_and_engineer_overrides(client):
    # cpu: +166% -> IMPROVED/verde (sugerido); trafico: +166% -> IMPROVED
    cfg = {"compare_days": ["2026-10-14", "2026-10-22"]}
    rows = {r["metric"]: r for r in post(client, "/api/analyze", SIMPLE, cfg).json["comparison"]["rows"]}
    assert (rows["NetSF cpu"]["status"], rows["NetSF cpu"]["color"]) == ("IMPROVED", "green")
    assert rows["NetSF cpu"]["auto_status"] == "IMPROVED"
    # El ingeniero pasa cpu a STABLE (el color sigue al estado) y pinta trafico de amarillo sin cambiar el estado
    cfg |= {"statuses": {"NetSF cpu": "STABLE", "bogus": "WHATEVER"},
            "var_colors": {"NetSF trafico": "yellow", "NetSF cpu": "purple"}}
    rows = {r["metric"]: r for r in post(client, "/api/analyze", SIMPLE, cfg).json["comparison"]["rows"]}
    assert (rows["NetSF cpu"]["status"], rows["NetSF cpu"]["color"]) == ("STABLE", "none")
    assert (rows["NetSF trafico"]["status"], rows["NetSF trafico"]["color"]) == ("IMPROVED", "yellow")
    ws = load_workbook(io.BytesIO(post(client, "/api/report/xlsx", SIMPLE, cfg).data)).active
    assert [ws.cell(row=4, column=k).value for k in (1, 2, 6, 7)] == ["Category", "KPI", "Var %", "Status"]
    assert ws.cell(row=4, column=3).value.startswith("Day 1")                # sin columna Group
    cpu, traf = (5, 6) if ws.cell(row=5, column=2).value == "NetSF cpu" else (6, 5)
    assert ws.cell(row=cpu, column=7).value == "STABLE" and ws.cell(row=cpu, column=6).fill.fgColor.rgb == "00000000"
    assert ws.cell(row=traf, column=6).fill.fgColor.rgb.endswith("FFF3BF")   # amarillo elegido
    assert ws.cell(row=traf, column=7).fill.fgColor.rgb.endswith("D3F9D8")   # IMPROVED = verde
    assert ws.cell(row=traf, column=1).value == "Other KPIs"   # "trafico" (no "traffic")
    assert post(client, "/api/report/pdf", SIMPLE, cfg).status_code == 200


def test_real_export_suggested_statuses(client):
    data = (SAMPLES / "ejemplo_5g_traffic.csv").read_bytes()
    rows = {r["metric"]: r for r in post(client, "/api/analyze", data, {"compare_days": ["2026-10-17", "2026-10-25"]})
            .json["comparison"]["rows"]}
    assert rows["NetSF Integrity"]["status"] == "STABLE"                                  # 0%
    assert rows["NetSF Average User Number"]["status"] == "DEGRADED"                      # -9.66%, mayor es mejor
    assert rows["NetSF Average Uplink PDCP Packet Loss Rate(%)"]["status"] == "IMPROVED"  # -41.95%, mayor es peor
    assert rows["NetSF Maximum User Number"]["status"] == "IMPROVED"                      # +16.27%


def test_integrity_has_no_chart_but_can_be_in_table(client):
    data = (SAMPLES / "ejemplo_5g_traffic.csv").read_bytes()
    j = post(client, "/api/analyze", data, {"compare_days": ["2026-10-17", "2026-10-25"]}).json
    charts = [c["metric"] for c in j["charts"]]
    assert "NetSF Integrity" not in charts and len(charts) == 8
    assert "NetSF Integrity" in j["metrics"]
    assert "NetSF Integrity" in [r["metric"] for r in j["comparison"]["rows"]]


def test_table_kpi_selection_is_independent(client):
    cfg = {"compare_days": ["2026-10-14", "2026-10-22"], "metrics": ["NetSF cpu"], "table_metrics": ["NetSF trafico"]}
    j = post(client, "/api/analyze", SIMPLE, cfg).json
    assert [c["metric"] for c in j["charts"]] == ["NetSF cpu"]
    assert [r["metric"] for r in j["comparison"]["rows"]] == ["NetSF trafico"]
    ws = load_workbook(io.BytesIO(post(client, "/api/report/xlsx", SIMPLE, cfg).data)).active
    assert ws.cell(row=5, column=2).value == "NetSF trafico" and ws.cell(row=6, column=2).value is None
    assert post(client, "/api/report/pdf", SIMPLE, cfg | {"table_metrics": []}).status_code == 200   # sin tabla


# ---------------------------------------------------------------- orden y categorías de KPI (4G y 5G)
from app.kpis import category, display_name, higher_is_better, sort_kpis  # noqa: E402

FIVE_G = [  # nombres tal como los envió operaciones, en desorden a propósito
    "NetSF N.NsaDc.SgNB.Rel", "NetSF Inter-SgNB PSCell Change SR(%)", "NetSF Cell CQI Avg",
    "NetSF SgNB Abnormal Release Rate (CU)(%)", "NetSF NR User UL Avg Throughput Mbps(Mbit/s)",
    "NetSF availability_netsf", "NetSF Sg-NB Addition Success Rate(%)", "NetSF Average User Number",
    "NetSF NR User DL Avg Throughput Mbps(Mbit/s)", "NetSF DL Traffic Volume GB(GB)", "NetSF CB RA Success Rate(%)",
    "NetSF Non-CB RA Success Rate(%)", "NetSF SgNB-Trig SgNB Abnorm Rel Rate(%)", "NetSF UL Traffic Volume GB(GB)",
    "NetSF Intra-SgNB IntraFreq PSCell Change Success Rate (CU)(%)", "NetSF UL Resource Block Utilizing Rate(%)",
    "NetSF N.User.NsaDc.PSCell.Avg", "NetSF Average Cell DL throughput Mbps(Mbit/s)",
]


def test_availability_netsf_display_name():
    assert display_name("availability_netsf") == "NetSF AVAILABILITY"
    assert display_name("NetSF availability_netsf") == "NetSF AVAILABILITY"
    assert display_name("Radio Network Availability Rate(%)") == "NetSF Radio Network Availability Rate(%)"
    assert display_name("NetSF eRAB DR(%)") == "NetSF eRAB DR(%)"


def test_5g_kpi_order():
    names = [display_name(k) for k in FIVE_G]
    out = sort_kpis(names)
    assert [category(k) for k in out] == sorted([category(k) for k in out], key=[
        "traffic", "users", "user_tput", "access", "retain", "mobility", "other", "counters"].index)
    assert out[:5] == ["NetSF DL Traffic Volume GB(GB)", "NetSF UL Traffic Volume GB(GB)", "NetSF Average User Number",
                       "NetSF NR User DL Avg Throughput Mbps(Mbit/s)", "NetSF NR User UL Avg Throughput Mbps(Mbit/s)"]
    assert category("NetSF AVAILABILITY") == "access" and category("NetSF Sg-NB Addition Success Rate(%)") == "access"
    assert category("NetSF CB RA Success Rate(%)") == "access" and category("NetSF Non-CB RA Success Rate(%)") == "access"
    assert category("NetSF SgNB Abnormal Release Rate (CU)(%)") == "retain"
    assert category("NetSF SgNB-Trig SgNB Abnorm Rel Rate(%)") == "retain"
    assert category("NetSF Inter-SgNB PSCell Change SR(%)") == "mobility"
    assert category("NetSF Intra-SgNB IntraFreq PSCell Change Success Rate (CU)(%)") == "mobility"
    assert category("NetSF N.User.NsaDc.PSCell.Avg") == "counters" and out[-2:] == [
        "NetSF N.NsaDc.SgNB.Rel", "NetSF N.User.NsaDc.PSCell.Avg"]


def test_4g_report_8_files_order_and_categories(client):
    files = sorted((SAMPLES / "4g").glob("*.csv"))
    assert len(files) == 8
    j = post(client, "/api/analyze", [(f.name, f.read_bytes()) for f in files],
             {"compare_days": ["2026-10-17", "2026-10-25"]}).json
    m = j["metrics"]
    assert len(m) == 62 and j["summary"]["rows"] == 30
    assert m[:3] == ["NetSF DL Data Volume(GB)(GB)", "NetSF UL data volume(GB)(GB)", "NetSF VoLTE Traffic(Erl)"]
    cats = [category(k) for k in m]
    order = ["traffic", "users", "user_tput", "access", "retain", "mobility", "other", "counters"]
    assert cats == sorted(cats, key=order.index)                                   # bloques en orden de prioridad
    ut = [k for k in m if category(k) == "user_tput"]
    assert ut == ["NetSF Average DL User Throughput(Mbit/s)", "NetSF Average UL User Throughput(Mbit/s)"]
    assert m.index("NetSF Radio Network Availability Rate(%)") < m.index("NetSF RRC Setup SR(%)")
    assert category("NetSF eRAB Setup SR QCI1(%)") == "access" and category("NetSF eRAB DR QCI1(%)") == "retain"
    assert category("NetSF LTE Drop Call(%)") == "retain" and category("NetSF RRC Re-Estab SR(%)") == "retain"
    assert category("NetSF Inter RAT HO SR(%)") == "mobility" and category("NetSF SRVCC SR(%)") == "mobility"
    assert category("NetSF Average DL Cell Throughput(Mbit/s)") == "other"
    assert all(category(k) == "counters" for k in m[-8:])                          # L.xxx al final
    assert not higher_is_better("NetSF eRAB DR(%)") and not higher_is_better("NetSF LTE Drop Call(%)")
    assert higher_is_better("NetSF RRC Setup SR(%)") and higher_is_better("NetSF Average DL User Throughput(Mbit/s)")
    assert [c["metric"] for c in j["charts"]] == [k for k in m if k != "NetSF Integrity"]
    assert j["charts"][0]["category"] == "traffic"
    assert [r["metric"] for r in j["comparison"]["rows"]] == m
    assert post(client, "/api/report/pdf", [(f.name, f.read_bytes()) for f in files],
                {"compare_days": ["2026-10-17", "2026-10-25"]}).status_code == 200
