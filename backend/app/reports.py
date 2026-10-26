"""Generación del PDF (gráficas + tablas) y del Excel de la tabla comparativa."""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4, landscape  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,  # noqa: E402
                                Spacer, Table, TableStyle)

from .analysis import AGGS, Config, _day_agg, chart_metrics, comparison, iter_series, summary
from .theme import (BORDER, DAY1_COLOR, DAY2_COLOR, INK, LABEL, MUTED, PRIMARY,
                    SERIES_COLORS, SOFT, VIOLET)
from .parsing import Dataset

GRID = "#ECEAF4"
from .analysis import STATUS_COLOR  # noqa: E402
from .kpis import CATEGORY_LABEL, category  # noqa: E402
from .theme import VAR_PALETTE  # noqa: E402


def fmt(v, pct=False):
    if v is None:
        return "—"
    s = f"{v:,.2f}"  # formato en inglés: 1,234.56
    return (("+" if v > 0 else "") + s + "%") if pct else s


# --------------------------------------------------------------------------- charts
def _short_num(v, *_):
    """Formato compacto como en la web (Intl compact en-US): 5.2M · 773.5K · 98"""
    a = abs(v)
    for div, suf in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if a >= div and (div > 1e3 or a >= 1e4):
            return f"{v / div:.1f}".removesuffix(".0") + suf
    return fmt(v).removesuffix(".00")


def render_chart(ds: Dataset, cfg: Config, metric: str, *, width=11.0, height=3.3, dpi=150) -> bytes:
    fig, ax = plt.subplots(figsize=(width, height), dpi=dpi)
    series = list(iter_series(ds, metric))
    for i, (name, part) in enumerate(series):
        ax.plot(part["__ts"], part[metric], lw=1.4, color=SERIES_COLORS[i % len(SERIES_COLORS)],
                marker="o" if len(part) <= 90 else None, ms=3,
                )

    # Días comparativos: banda + etiqueta con el valor del KPI (solo en estos dos días)
    # Con datos diarios cada día es un solo punto (a las 00:00): las bandas se centran en ese punto
    daily = (summary(ds).get("step_minutes") or 0) >= 1440
    legend = []
    shift = timedelta(hours=12) if daily else timedelta(0)
    if cfg.compare_days:
        for k, (day, color) in enumerate(zip(cfg.compare_days, (DAY1_COLOR, DAY2_COLOR))):
            d0 = datetime.combine(day, datetime.min.time())
            ax.axvspan(d0 - shift, d0 + timedelta(days=1) - shift, color=color, alpha=0.10, lw=0)
            x = d0 if daily else d0 + timedelta(hours=12)
            for i, (name, part) in enumerate(series):
                v, _ = _day_agg(part, metric, day, cfg.agg)
                if v is None:
                    continue
                ax.annotate(_short_num(v), xy=(x, v), xytext=(0, 9 + 11 * (i % 3)), textcoords="offset points",
                            ha="center", fontsize=7.5, fontweight="bold", color="white",
                            bbox=dict(boxstyle="round,pad=0.25", fc=color, ec="none"))
                ax.plot([x], [v], "o", ms=5, color=color, zorder=5)
            legend.append(Patch(facecolor=color, alpha=0.5, label=f"Day {k + 1} · {day:%d/%m}"))

    for i, e in enumerate(cfg.executions):
        if e.mode == "day":
            ax.axvspan(e.start - shift, e.end - shift, color=e.color, alpha=0.08, lw=0)
        ax.axvline(e.start, color=e.color, lw=2, ls="--")
        legend.append(Line2D([0], [0], color=e.color, lw=2, ls="--", label=e.label))
        ax.annotate(e.label, xy=(e.start, 1), xycoords=("data", "axes fraction"), xytext=(4, -4 - (i % 4) * 11),
                    textcoords="offset points", va="top", fontsize=8, color=e.color, fontweight="bold",
                    bbox=dict(boxstyle="square,pad=0.1", fc="white", ec="none", alpha=0.8))
    ax.margins(y=0.18)

    ax.set_title(metric, loc="left", fontsize=12, color=INK, fontweight="bold")
    ax.grid(True, color=GRID, lw=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(BORDER)
    ax.tick_params(colors=MUTED, labelsize=8)
    ax.yaxis.set_major_formatter(FuncFormatter(_short_num))
    loc = mdates.AutoDateLocator()
    ax.xaxis.set_major_locator(loc)
    ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(loc))
    if legend:
        ax.legend(handles=legend, fontsize=7.5, frameon=False, ncol=min(8, len(legend)),
                  loc="lower right", bbox_to_anchor=(1, 1.0), borderaxespad=0.2)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF
def build_pdf(ds: Dataset, cfg: Config, filename: str) -> bytes:
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=cfg.title)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], alignment=0, fontSize=20, textColor=colors.HexColor(INK))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor(INK))
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, textColor=colors.HexColor(MUTED))
    kicker = ParagraphStyle("k", parent=body, fontSize=8, textColor=colors.HexColor(LABEL), fontName="Helvetica-Bold")
    small = ParagraphStyle("s", parent=body, fontSize=8)
    cell = ParagraphStyle("c", parent=body, fontSize=8.5, textColor=colors.HexColor(INK))

    info = summary(ds)
    daily = (info.get("step_minutes") or 0) % 1440 == 0
    story = [Paragraph("NETSF · NETWORK OPERATIONS", kicker), Paragraph(cfg.title, h1),
             Paragraph(f"Source: <b>{_esc(_short(filename, 90))}</b> · {info['rows']:,} rows · "
                       f"{_d(info['start'], daily)} → {_d(info['end'], daily)} · generated {datetime.now():%d/%m/%Y %H:%M}", body),
             Spacer(1, 8)]
    marks = [("", e.label, e.describe(), e.color) for e in cfg.executions]
    if cfg.compare_days:
        marks += [("", f"Comparison day {k + 1}", d.strftime("%d/%m/%Y") + " (full day)", c)
                  for k, (d, c) in enumerate(zip(cfg.compare_days, (DAY1_COLOR, DAY2_COLOR)))]
    if marks:
        t = Table([["", "Marker", "Date"]] + [list(m[:3]) for m in marks], colWidths=[0.5 * cm, 8 * cm, 7 * cm],
                  hAlign="LEFT")
        style = _table_style()
        for i, m in enumerate(marks, start=1):
            style.add("BACKGROUND", (0, i), (0, i), colors.HexColor(m[3]))
        t.setStyle(style)
        story += [t, Spacer(1, 10)]

    avail_w = page[0] - doc.leftMargin - doc.rightMargin
    section = ParagraphStyle("sec", parent=h2, fontSize=12, textColor=colors.HexColor(PRIMARY), spaceBefore=4,
                             spaceAfter=4)
    last_cat = None
    for m in chart_metrics(ds, cfg):   # ya vienen en orden de prioridad (kpis.sort_kpis)
        png = render_chart(ds, cfg, m)
        img = Image(io.BytesIO(png), width=avail_w, height=avail_w * 3.3 / 11.0)
        block = [img]
        if category(m) != last_cat:     # encabezado de sección: Traffic, Users, ...
            last_cat = category(m)
            block.insert(0, Paragraph(CATEGORY_LABEL[last_cat], section))
        story += [KeepTogether(block), Spacer(1, 6)]

    c = comparison(ds, cfg)
    if c and c["rows"]:   # sin KPI seleccionados para la tabla -> no se incluye
        story.append(PageBreak())
        story.append(Paragraph(f"Comparison table — full-day {c['agg_label'].lower()}", h2))
        story.append(Paragraph(
            f"<font color='{DAY1_COLOR}'>■</font> Day 1: <b>{_d(c['day1'])}</b> &nbsp; vs &nbsp; "
            f"<font color='{DAY2_COLOR}'>■</font> Day 2: <b>{_d(c['day2'])}</b> · Difference = Day 2 - Day 1 · "
            f"Status: STABLE within ±{c['stable_pct']:g}%, otherwise IMPROVED / DEGRADED (reviewed by engineering)", body))
        data = [["KPI"] +
                [f"Day 1\n{_d(c['day1'])}", f"Day 2\n{_d(c['day2'])}", "Difference", "Var %", "Status"]]
        row_of, sections, last_cat = [], [], None   # fila de la tabla de cada KPI / filas de sección
        for r in c["rows"]:
            if r["category"] != last_cat:
                last_cat = r["category"]
                sections.append(len(data))
                data.append([CATEGORY_LABEL[last_cat], "", "", "", "", ""])
            row_of.append(len(data))
            data.append([Paragraph(_esc(r["metric"]), cell)] +
                        [fmt(r["v1"]), fmt(r["v2"]), fmt(r["delta"]), fmt(r["delta_pct"], pct=True), r["status"] or "—"])
        num_w = 3.6 * cm
        kpi_w = (page[0] - doc.leftMargin - doc.rightMargin) - 5 * num_w
        t = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[kpi_w] + [num_w] * 5)
        st = _table_style(header_bg=True)
        first_num = 1
        st.add("ALIGN", (first_num, 0), (-2, -1), "RIGHT")
        st.add("ALIGN", (-1, 0), (-1, -1), "CENTER")
        st.add("TEXTCOLOR", (first_num, 0), (first_num, 0), colors.HexColor(DAY1_COLOR))
        st.add("TEXTCOLOR", (first_num + 1, 0), (first_num + 1, 0), colors.HexColor(DAY2_COLOR))
        for i in sections:
            st.add("SPAN", (0, i), (-1, i))
            st.add("BACKGROUND", (0, i), (-1, i), colors.HexColor(SOFT))
            st.add("FONT", (0, i), (-1, i), "Helvetica-Bold", 8.5)
            st.add("TEXTCOLOR", (0, i), (-1, i), colors.HexColor(PRIMARY))
            st.add("ALIGN", (0, i), (-1, i), "LEFT")
        for i, r in zip(row_of, c["rows"]):
            # Var % (col -2) y Status (col -1) con los colores definidos/revisados por el ingeniero
            for col, key in ((-2, r.get("color")), (-1, STATUS_COLOR.get(r.get("status") or "", "none"))):
                if key in VAR_PALETTE:
                    bg, ink = VAR_PALETTE[key]
                    st.add("BACKGROUND", (col, i), (col, i), colors.HexColor(bg))
                    st.add("TEXTCOLOR", (col, i), (col, i), colors.HexColor(ink))
                    st.add("FONT", (col, i), (col, i), "Helvetica-Bold", 8.5)
            if r.get("status"):
                st.add("FONT", (-1, i), (-1, i), "Helvetica-Bold", 8)
        t.setStyle(st)
        story += [Spacer(1, 6), t]
        if c["missing_day1"]:
            story.append(Paragraph("No data for Day 1.", small))
        if c["missing_day2"]:
            story.append(Paragraph("No data for Day 2.", small))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawRightString(page[0] - d.rightMargin, 0.6 * cm, f"{cfg.title} · page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _table_style(header_bg=True) -> TableStyle:
    return TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 8.5),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(SOFT if header_bg else "#ffffff")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(INK)),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor(BORDER)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ])


def _d(iso, date_only=False):
    if not iso:
        return "—"
    return datetime.fromisoformat(iso).strftime("%d/%m/%Y" if date_only or len(iso) == 10 else "%d/%m/%Y %H:%M")


def _short(name: str, max_len: int = 45) -> str:  # noqa: D401
    """Quita marcas de tiempo largas (p. ej. _20261026105118484) y recorta nombres muy largos."""
    s = re.sub(r"_?\d{8,}", "", name)
    return s if len(s) <= max_len else s[: max_len - 12] + "…" + s[-11:]


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_page_images(pdf: bytes, scale: float = 1.4, max_pages: int = 60) -> list[bytes]:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(pdf)
    out = []
    try:
        for i in range(min(len(doc), max_pages)):
            img = doc[i].render(scale=scale).to_pil()
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG", optimize=True)
            out.append(buf.getvalue())
    finally:
        doc.close()
    return out


# --------------------------------------------------------------------------- XLSX
def build_xlsx(ds: Dataset, cfg: Config) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparativa"
    hex_ = lambda c: c.lstrip("#").upper()  # noqa: E731
    bold = Font(bold=True, color=hex_(INK))
    head_fill = PatternFill("solid", fgColor=hex_(SOFT))
    thin = Border(bottom=Side(style="thin", color=hex_(BORDER)))
    num = '#,##0.00'
    c = comparison(ds, cfg)

    ws.append([cfg.title])
    ws["A1"].font = Font(bold=True, size=14, color=hex_(PRIMARY))
    if c:
        ws.append([f"Full-day {AGGS[cfg.agg].lower()} · Day 1: {_d(c['day1'])} · Day 2: {_d(c['day2'])} · "
                   "Difference = Day 2 - Day 1"])
    else:
        ws.append(["Select both comparison days to build the table."])
    ws.append([])

    headers = ["Category", "KPI"] + \
              [f"Day 1 ({_d(c['day1'])})" if c else "Day 1", f"Day 2 ({_d(c['day2'])})" if c else "Day 2",
               "Difference", "Var %", "Status", "Samples Day 1", "Samples Day 2"]
    ws.append(headers)
    hr = ws.max_row
    for cell in ws[hr]:
        cell.font, cell.fill, cell.border = bold, head_fill, thin
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    off = 1   # columna Category al inicio
    ws.cell(row=hr, column=2 + off).font = Font(bold=True, color=hex_(DAY1_COLOR))
    ws.cell(row=hr, column=3 + off).font = Font(bold=True, color=hex_(DAY2_COLOR))

    for r in (c["rows"] if c else []):
        ws.append([CATEGORY_LABEL[r["category"]], r["metric"]] +
                  [r["v1"], r["v2"], r["delta"], None if r["delta_pct"] is None else r["delta_pct"] / 100,
                   r["status"], r["n1"], r["n2"]])
        row = ws[ws.max_row]
        for k in (1, 2, 3):
            row[k + off].number_format = num
        row[4 + off].number_format = '+0.00%;-0.00%;0.00%'
        row[5 + off].alignment = Alignment(horizontal="center")
        for k, key in ((4, r.get("color")), (5, STATUS_COLOR.get(r.get("status") or "", "none"))):
            if key in VAR_PALETTE:   # colores definidos/revisados por el ingeniero
                bg, ink = VAR_PALETTE[key]
                row[k + off].font = Font(bold=True, color=hex_(ink))
                row[k + off].fill = PatternFill("solid", fgColor=hex_(bg))

    widths = [26, 44, 18, 18, 16, 10, 12, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{max(ws.max_row, hr)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
